"""文本抽取：把各种格式的附件解析成正文，以及上传前的图片压缩。"""

import csv
import io
import os

from attachment.config import (
    AttachmentError,
    IMAGE_MAX_EDGE,
    MAX_IMAGE_BYTES,
    MAX_SHEET_COLS,
    MAX_SHEET_ROWS,
)


# ---------------- 文本解码 / 表格转 Markdown ----------------
def _decode(data: bytes) -> str:
    """按常见中文编码依次尝试解码，最后兜底替换非法字节。"""
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _rows_to_markdown(rows: list) -> str:
    """把二维表转成 Markdown 表格，首行当表头。

    注意：openpyxl 按固定列宽取数会产生大量尾部空单元格，这里统一裁掉，
    否则表格会多出一堆空列，白白吃掉上下文。
    """
    def trim(row: list) -> list:
        cells = [("" if c is None else str(c)) for c in row]
        end = len(cells)
        while end > 0 and not cells[end - 1].strip():
            end -= 1
        return cells[:end]

    rows = [trim(r) for r in rows]
    rows = [r for r in rows if r]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]

    def esc(value: str) -> str:
        return value.replace("|", "\\|").replace("\r", " ").replace("\n", " ").strip()

    lines = [
        "| " + " | ".join(esc(c) for c in rows[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(esc(c) for c in r) + " |" for r in rows[1:])
    return "\n".join(lines)


def _parse_csv(data: bytes) -> str:
    text = _decode(data)
    rows = list(csv.reader(io.StringIO(text)))
    truncated = len(rows) > MAX_SHEET_ROWS
    rows = rows[:MAX_SHEET_ROWS]
    table = _rows_to_markdown([r[:MAX_SHEET_COLS] for r in rows])
    if truncated:
        table += f"\n（表格行数过多，仅解析前 {MAX_SHEET_ROWS} 行）"
    return table


def _parse_xlsx(path: str) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # 依赖缺失时给出可执行的提示
        raise AttachmentError("服务端未安装 openpyxl，无法解析 Excel（pip install openpyxl）") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            rows = [
                list(row)
                for row in ws.iter_rows(
                    max_row=MAX_SHEET_ROWS + 1, max_col=MAX_SHEET_COLS, values_only=True
                )
            ]
            table = _rows_to_markdown(rows[:MAX_SHEET_ROWS])
            if not table:
                continue
            head = f"### 工作表：{ws.title}"
            if len(rows) > MAX_SHEET_ROWS:
                head += f"（仅前 {MAX_SHEET_ROWS} 行 × {MAX_SHEET_COLS} 列）"
            parts.append(f"{head}\n{table}")
    finally:
        wb.close()
    return "\n\n".join(parts)


def _parse_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise AttachmentError("服务端未安装 pypdf，无法解析 PDF（pip install pypdf）") from exc

    reader = PdfReader(path)
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # 单页损坏不影响整篇
            text = ""
        if text.strip():
            pages.append(f"### 第 {index} 页\n{text.strip()}")
    if not pages:
        raise AttachmentError("该 PDF 未解析出文本（可能是扫描件/纯图片），请改用图片或提供文字版")
    return "\n\n".join(pages)


def _parse_docx(path: str) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise AttachmentError("服务端未安装 python-docx，无法解析 Word（pip install python-docx）") from exc

    document = Document(path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        markdown = _rows_to_markdown(rows)
        if markdown:
            parts.append(f"### 表格 {index}\n{markdown}")
    return "\n".join(parts)


def extract_text(kind: str, path: str, data: bytes) -> str:
    """按类型抽取正文；图片不需要解析，直接返回空串。"""
    if kind == "image":
        return ""
    if kind == "pdf":
        return _parse_pdf(path)
    if kind == "word":
        return _parse_docx(path)
    if kind == "sheet":
        if path.lower().endswith(".csv"):
            return _parse_csv(data)
        return _parse_xlsx(path)
    if kind == "text":
        return _decode(data)
    return ""


# ---------------- 图片压缩（可选依赖 Pillow） ----------------
def shrink_image_if_needed(path: str, ext: str, size: int) -> tuple:
    """把超大/像素过多的图片压到 IMAGE_MAX_EDGE 以内，返回 (path, ext, size)。

    没装 Pillow、或（既没超像素也没超体积）时原样返回，不影响主流程。
    """
    if ext == "gif":
        return path, ext, size
    try:
        from PIL import Image
    except ImportError:
        return path, ext, size

    try:
        with Image.open(path) as img:
            width, height = img.size
            longest = max(width, height)
            need_resize = longest > IMAGE_MAX_EDGE
            if not need_resize and size <= MAX_IMAGE_BYTES // 2:
                return path, ext, size
            if need_resize:
                scale = IMAGE_MAX_EDGE / float(longest)
                img = img.resize(
                    (max(1, int(width * scale)), max(1, int(height * scale)))
                )
            buffer = io.BytesIO()
            img.convert("RGB").save(buffer, format="JPEG", quality=85, optimize=True)
    except Exception:  # 压缩失败就用原图，不能因此让上传失败
        return path, ext, size

    new_data = buffer.getvalue()
    # 只是体积偏大时，压缩没变小就保留原图；像素超限时则以"缩到上限"为准
    if len(new_data) >= size and not need_resize:
        return path, ext, size

    new_path = os.path.splitext(path)[0] + ".jpg"
    with open(new_path, "wb") as handle:
        handle.write(new_data)
    if new_path != path:
        try:
            os.remove(path)
        except OSError:
            pass
    return new_path, "jpg", len(new_data)
